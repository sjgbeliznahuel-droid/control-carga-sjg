import datetime
import io
import os
import re
import time
import unicodedata
import json
import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

st.set_page_config(
    page_title="Control de Fichadas - SJG", page_icon="🏭", layout="wide"
)

# Inicializar estado de sesión
if "mostrar_reporte" not in st.session_state:
    st.session_state["mostrar_reporte"] = False
if "cruce_cierre_ejecutado" not in st.session_state:
    st.session_state["cruce_cierre_ejecutado"] = False

# --- CONFIGURACIÓN GOOGLE SHEETS & DRIVE ---
SHEET_ID = "1gGkGjO-ZX7auYdlKAuI9rLEN_KoVah3vdnz7WofcHq0"
URL_DESCARGA_SHEET = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
WEBAPP_URL = "https://script.google.com/macros/s/AKfycbxZOKYHQHzLR-gZionm3Ub1z8Rlrz5l6AtR1D8gLKKIeAqhSGMf62vMfsa4F2s1gcweFA/exec"
DRIVE_FOLDER_ID = "1fMm6XYmw-dwqSqhijX6tHAe2VdVMtA2-"
NOTIF_FOLDER_ID = "1GJGTwWRwmrWT-Zgkp4nQcWmm3kk-eR8y"  # Carpeta de Notificaciones Finales

# --- CONEXIÓN A GOOGLE DRIVE ---
@st.cache_resource
def conectar_drive():
    credenciales_dict = json.loads(st.secrets["credenciales"])
    creds = service_account.Credentials.from_service_account_info(
        credenciales_dict, scopes=['https://www.googleapis.com/auth/drive']
    )
    return build('drive', 'v3', credentials=creds)

drive_service = conectar_drive()

def normalizar_ot(val):
    """Convierte cualquier valor de OT a un string numérico limpio sin decimales ni espacios"""
    if pd.isna(val): return ""
    s = str(val).strip()
    if not s or s.lower() in ["nan", "none", "null"]: return ""
    s = re.sub(r'\.0+$', '', s)
    try:
        val_float = float(s)
        if val_float.is_integer():
            return str(int(val_float))
    except Exception:
        pass
    return s

def obtener_archivos_drive():
    query = f"'{DRIVE_FOLDER_ID}' in parents and trashed=false"
    resultados = drive_service.files().list(q=query, orderBy="name desc", pageSize=1000, fields="files(id, name)").execute()
    return resultados.get('files', [])

def obtener_archivos_notif_drive():
    """Lista todos los archivos (Excel o Google Sheets) de la carpeta de notificaciones en Drive"""
    try:
        query = f"'{NOTIF_FOLDER_ID}' in parents and trashed=false"
        resultados = drive_service.files().list(
            q=query, 
            orderBy="modifiedTime desc", 
            pageSize=1000, 
            fields="files(id, name, mimeType)"
        ).execute()
        
        archivos = resultados.get('files', [])
        archivos_validos = []
        for f in archivos:
            nombre = f.get('name', '')
            mime = f.get('mimeType', '')
            if nombre.lower().endswith(('.xlsx', '.xltx', '.csv')) or mime == 'application/vnd.google-apps.spreadsheet':
                archivos_validos.append(f)
                
        return archivos_validos
    except Exception as e:
        st.error(f"⚠️ Error de acceso a la carpeta de Drive: {e}")
        st.info("💡 Asegurate de haber compartido la carpeta con el mail de tu cuenta de servicio.")
        return []

def descargar_archivo_drive(file_id):
    request = drive_service.files().get_media(fileId=file_id)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    terminado = False
    while not terminado:
        _, terminado = downloader.next_chunk()
    fh.seek(0)
    return fh

def descargar_archivo_notif_drive(file_id, mime_type):
    """Descarga tanto archivos Excel físicos como Hojas de Cálculo nativas de Google"""
    if mime_type == 'application/vnd.google-apps.spreadsheet':
        request = drive_service.files().export_media(fileId=file_id, mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    else:
        request = drive_service.files().get_media(fileId=file_id)
        
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    terminado = False
    while not terminado:
        _, terminado = downloader.next_chunk()
    fh.seek(0)
    return fh

def obtener_notificaciones_drive(file_id, mime_type=None):
    """Lee la planilla de notificaciones seleccionada detectando automáticamente columnas por patrón de contenido"""
    if not file_id:
        return {}, {}
    try:
        fh = descargar_archivo_notif_drive(file_id, mime_type)
        df_raw = pd.read_excel(fh, header=None)
        
        # Buscar fila de encabezados en las primeras 15 filas
        header_row = 0
        for idx in range(min(15, len(df_raw))):
            row_vals = [str(v).upper() for v in df_raw.iloc[idx].values if pd.notna(v)]
            row_str = " ".join(row_vals)
            if any(k in row_str for k in ["OT", "SOLICITUD", "ORDEN", "NOTIF"]):
                header_row = idx
                break
        
        raw_headers = [str(v).strip() if pd.notna(v) else f"Col_{i}" for i, v in enumerate(df_raw.iloc[header_row].values)]
        df_notif = df_raw.iloc[header_row + 1:].copy()
        df_notif.columns = raw_headers
        
        # 1. Identificar columna OT
        ot_col = None
        for c in df_notif.columns:
            c_str = str(c).strip().upper()
            if c_str in ["OT", "ORDEN", "N° OT", "Nº OT", "N° ORDEN", "ORDEN DE TRABAJO", "OT PROGRAMADA"]:
                ot_col = c
                break
        if not ot_col:
            for c in df_notif.columns:
                c_str = str(c).strip().upper()
                if re.search(r'\b(OT|ORDEN)\b', c_str):
                    ot_col = c
                    break
        if not ot_col:
            for c in df_notif.columns:
                sample = [normalizar_ot(v) for v in df_notif[c].dropna().head(20)]
                if any(len(v) >= 7 and v.isdigit() for v in sample):
                    ot_col = c
                    break

        # 2. Identificar columna Solicitud
        sol_col = None
        for c in df_notif.columns:
            c_str = str(c).strip().upper()
            if any(k in c_str for k in ["SOLICITUD", "SOLIC", "N° SOL", "NUMERO SOL", "NRO SOL"]):
                sol_col = c
                break

        # 3. Identificar columna Estado
        est_col = None
        for c in df_notif.columns:
            c_str = str(c).strip().upper()
            if any(k in c_str for k in ["ESTADO", "STATUS"]):
                est_col = c
                break

        # 4. Identificar columna Fecha (Columna D = Índice 3 por defecto)
        fecha_col = None
        if len(df_notif.columns) > 3:
            fecha_col = df_notif.columns[3]
        else:
            for c in df_notif.columns:
                c_str = str(c).strip().upper()
                if any(k in c_str for k in ["FECHA", "NOTIF"]):
                    fecha_col = c
                    break

        mapa_notif_acumulado = {}
        if ot_col:
            for _, row in df_notif.iterrows():
                ot_key = normalizar_ot(row[ot_col])
                if not ot_key: continue

                val_sol = str(row[sol_col]).strip() if sol_col and pd.notna(row[sol_col]) and str(row[sol_col]).lower() != "nan" else ""
                val_est = str(row[est_col]).strip() if est_col and pd.notna(row[est_col]) and str(row[est_col]).lower() != "nan" else ""

                val_fecha = ""
                if fecha_col and pd.notna(row[fecha_col]) and str(row[fecha_col]).lower() != "nan":
                    v_f = row[fecha_col]
                    if isinstance(v_f, (pd.Timestamp, datetime.date, datetime.datetime)):
                        val_fecha = v_f.strftime("%d/%m/%Y")
                    else:
                        str_vf = str(v_f).strip()
                        match_dt = re.search(r'(\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4})', str_vf)
                        if match_dt:
                            val_fecha = match_dt.group(1)
                        else:
                            val_fecha = str_vf.split()[0] if " " in str_vf else str_vf

                if ot_key not in mapa_notif_acumulado:
                    mapa_notif_acumulado[ot_key] = {
                        "solicitudes": [],
                        "notificaciones": [],
                        "estados": []
                    }

                if val_sol and val_sol not in mapa_notif_acumulado[ot_key]["solicitudes"]:
                    mapa_notif_acumulado[ot_key]["solicitudes"].append(val_sol)
                if val_fecha and val_fecha not in mapa_notif_acumulado[ot_key]["notificaciones"]:
                    mapa_notif_acumulado[ot_key]["notificaciones"].append(val_fecha)
                if val_est and val_est not in mapa_notif_acumulado[ot_key]["estados"]:
                    mapa_notif_acumulado[ot_key]["estados"].append(val_est)

        mapa_notif_final = {}
        for k_ot, v_data in mapa_notif_acumulado.items():
            mapa_notif_final[k_ot] = {
                "solicitud": " - ".join(v_data["solicitudes"]),
                "notificacion": " - ".join(v_data["notificaciones"]),
                "estado": " - ".join(v_data["estados"])
            }

        diag = {
            "ot_col": ot_col,
            "sol_col": sol_col,
            "est_col": est_col,
            "fecha_col": fecha_col,
            "total_keys": len(mapa_notif_final)
        }

        return mapa_notif_final, diag
    except Exception as e:
        st.error(f"Error al procesar la planilla de notificaciones: {e}")
        return {}, {}

# --- LOGO VECTORIAL ---
SJG_LOGO_SVG = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 420 130" width="100%" height="75" style="max-width: 380px;"><g fill="#0A0A0A"><text x="5" y="82" font-family="\'Montserrat\', \'Arial Black\', sans-serif" font-weight="900" font-size="78" letter-spacing="-1.5">SJG</text><text x="8" y="116" font-family="\'Montserrat\', \'Arial\', sans-serif" font-weight="800" font-size="17" letter-spacing="3.2">MONTAJES INDUSTRIALES</text><polygon points="230,126 260,126 282,98 230,126" /><polygon points="318,24 294,126 268,126 290,98 318,24" /><polygon points="352,70 302,126 392,126" /></g></svg>'

col_logo, col_titulo = st.columns([1.2, 3])
with col_logo:
    st.markdown(SJG_LOGO_SVG, unsafe_allow_html=True)
with col_titulo:
    st.title("Sistema de Control y Validación de Personal")
    st.caption("Conciliación automatizada de Partes de Carga vs. Fichadas en Puerta (Cloud - En Vivo)")
st.divider()

MESES_NOMBRES = {1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre"}

EXCLUIDOS_SECTOR = {
    60003127: ("Villaseca Urrutia, Silvina", "Seguridad e Higiene"),
    60017753: ("Miño, Dario Oscar", "Seguridad e Higiene"),
    60016617: ("Ferreyra, Valentin", "Seguridad e Higiene"),
    50027446: ("Romano, Nicolas", "Delegados"),
    50012920: ("Chavez, Fernando", "Delegados"),
    50027313: ("Carrizo, Pedro Paz", "Delegados"),
    50012909: ("Andenmatten, Marcelo", "Delegados"),
    50029799: ("Sinner, Damian", "Personal de Pañol"),
    60008058: ("Nuñez, Gabriel", "Personal de Pañol"),
    60002019: ("Gonzalez, Hector", "Personal de Pañol"),
    60008120: ("Fourcade, Hernan", "Personal de Pañol"),
    60022214: ("Franco, Martin Jesus", "Personal de Pañol"),
    50034415: ("Brozman, Ariel", "Personal de Pañol"),
    60008497: ("Diaz Paez, Alejandro", "Coordinación"),
    60012942: ("Cometti, Juliana", "Coordinación"),
    60003811: ("Gauna, Mauricio", "Coordinación"),
    60008266: ("Lascano, Luis Alejandro", "Coordinación"),
}

def limpiar_legajo(val):
    if pd.isna(val): return None
    try: return int(float(str(val).strip()))
    except (ValueError, TypeError): return None

def obtener_excel_nube(url):
    response = requests.get(url)
    response.raise_for_status()
    return response.content

def normalizar_texto(val):
    if pd.isna(val): return ""
    s = str(val).lower().strip()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return " ".join("".join(c for c in s if c.isalnum() or c.isspace()).split())

def obtener_dia_mes(texto_o_fecha):
    if isinstance(texto_o_fecha, (pd.Timestamp, datetime.date, datetime.datetime)):
        return texto_o_fecha.day, texto_o_fecha.month
    match = re.search(r"(\d{1,2})[-/\.](\d{1,2})", str(texto_o_fecha))
    if match:
        try:
            d, m = int(match.group(1)), int(match.group(2))
            if 1 <= d <= 31 and 1 <= m <= 12: return d, m
        except Exception: pass
    return None, None

def extraer_fecha_contenido(df_raw):
    for idx in range(min(10, len(df_raw))):
        for val in df_raw.iloc[idx].dropna():
            if isinstance(val, (pd.Timestamp, datetime.date, datetime.datetime)):
                return val.day, val.month
            match = re.search(r"(\d{1,2})[-/\.](\d{1,2})[-/\.](\d{2,4})", str(val))
            if match:
                try:
                    d, m = int(match.group(1)), int(match.group(2))
                    if 1 <= d <= 31 and 1 <= m <= 12: return d, m
                except Exception: pass
    return None, None

def extraer_horas(val):
    if pd.isna(val): return 0.0
    match = re.search(r"(\d+(?:[\.,]\d+)?)", str(val).upper().strip())
    if match: return float(match.group(1).replace(",", "."))
    return 0.0

def calcular_horas_fichada(ing, sal):
    try:
        i, s = float(ing), float(sal)
        return s - i if s >= i else (24.0 - i) + s
    except Exception: return 0.0

t_rrhh, t_carga, t_reglas, t_reporte, t_cierre, t_quincena = st.tabs([
    "📥 Carga de Fichadas/RRHH", "📂 1. Carga", "⚙️ 2. Reglas", "📊 3. Reporte", "📋 Cierre de Período", "⏱️ 4. Horas y Quincenas"
])

with t_rrhh:
    st.header("Carga Diaria de Fichadas (Portería / RRHH)")
    st.info("ℹ️ Para cargar una nueva fichada, arrastrá el Excel exportado directamente a tu carpeta 'FICHADAS' en Google Drive.")
    st.success("La aplicación detectará automáticamente los nuevos archivos cuando los subas a Drive.")

with t_carga:
    c1, c2 = st.columns(2)
    with c2:
        st.subheader("Base de Fichadas Guardadas (Drive)")
        try:
            archivos_drive = obtener_archivos_drive()
            archivos_guardados = [f['name'] for f in archivos_drive if f['name'].endswith(('.xlsx', '.xltx'))]
            mapa_archivos = {f['name']: f['id'] for f in archivos_drive}
            
            if not archivos_guardados:
                st.warning("Aún no hay fichadas en Drive.")
                af_selected = None
            else:
                fichadas_por_mes = {}
                for f in archivos_guardados:
                    match = re.search(r"Fichada_(\d{1,2})-(\d{1,2})", f)
                    if match:
                        dia, mes = int(match.group(1)), int(match.group(2))
                        if mes not in fichadas_por_mes: fichadas_por_mes[mes] = {}
                        fichadas_por_mes[mes][dia] = f

                if not fichadas_por_mes:
                    opciones = ["-- Seleccione una fichada --"] + archivos_guardados
                    af_selected = st.selectbox("Elegir Fichada:", options=opciones)
                    if af_selected == "-- Seleccione una fichada --": af_selected = None
                else:
                    col_mes, col_dia = st.columns(2)
                    with col_mes:
                        meses_ordenados = sorted(fichadas_por_mes.keys(), reverse=True)
                        mes_elegido_num = st.selectbox("Mes de Fichada:", options=meses_ordenados, format_func=lambda m: f"📅 {MESES_NOMBRES.get(m, 'Mes')}")
                    with col_dia:
                        dias_del_mes = sorted(fichadas_por_mes[mes_elegido_num].keys(), reverse=True)
                        opciones_dia = ["-- Seleccionar Día --"] + [f"Día {d:02d}" for d in dias_del_mes]
                        dia_elegido_str = st.selectbox("Día de Fichada:", options=opciones_dia)

                    if dia_elegido_str != "-- Seleccionar Día --":
                        dia_elegido_num = int(dia_elegido_str.replace("Día ", ""))
                        af_selected = fichadas_por_mes[mes_elegido_num][dia_elegido_num]
                        st.info(f"☁️ Archivo vinculado desde Drive: **{af_selected}**")
                    else:
                        af_selected = None
        except Exception as e:
            st.error("Error al leer archivos de Drive.")
            af_selected = None

    with c1:
        st.subheader("Planilla de Horas (Carga)")
        st.success("☁️ Sincronizado con Google Sheets")
        try:
            with st.spinner("Sincronizando con Google Sheets..."):
                carga_bytes = obtener_excel_nube(URL_DESCARGA_SHEET)
                xl_carga = pd.ExcelFile(io.BytesIO(carga_bytes))
                hojas_disponibles = xl_carga.sheet_names
            nh = st.selectbox("Seleccionar Fecha / Pestaña:", options=hojas_disponibles)
        except Exception as e:
            st.error(f"No se pudo conectar a Google Sheets. Error: {e}")
            carga_bytes, nh = None, None

with t_reglas:
    st.subheader("⚙️ Configuración de Personal Excluido / No Facturable")
    df_reglas = pd.DataFrame([{"Legajo": leg, "Nombre": datos[0], "Sector / Área": datos[1]} for leg, datos in EXCLUIDOS_SECTOR.items()])
    st.dataframe(df_reglas, use_container_width=True)

with t_reporte:
    if carga_bytes and af_selected:
        col_btn1, col_btn2 = st.columns([1, 1])
        with col_btn1:
            if st.button("🚀 Ejecutar Cruce Estricto por Fecha", type="primary"):
                st.session_state["mostrar_reporte"] = True
                if hasattr(st, "rerun"): st.rerun()
                else: st.experimental_rerun()
        with col_btn2:
            if st.button("🔄 Sincronizar Todo (Drive y Sheets)"):
                st.session_state["mostrar_reporte"] = True
                if hasattr(st, "rerun"): st.rerun()
                else: st.experimental_rerun()

        if st.session_state["mostrar_reporte"]:
            try:
                af_selected_id = mapa_archivos.get(af_selected)
                with st.spinner("Descargando fichada desde Drive..."):
                    fh_fichada = descargar_archivo_drive(af_selected_id)

                d_pestana, m_pestana = obtener_dia_mes(nh)
                dff_raw = pd.read_excel(fh_fichada, sheet_name="Hoja2", header=None)
                fh_fichada.seek(0) 
                
                d_fichada, m_fichada = obtener_dia_mes(af_selected)
                if d_fichada is None: d_fichada, m_fichada = extraer_fecha_contenido(dff_raw)

                if d_pestana and d_fichada and (d_pestana, m_pestana) != (d_fichada, m_fichada):
                    st.error(f"⛔ **Error de Validación:** La Pestaña `{nh}` es del **{d_pestana:02d}/{m_pestana:02d}**, pero la Fichada es del **{d_fichada:02d}/{m_fichada:02d}**.")
                    st.stop()

                dfc_raw = pd.read_excel(io.BytesIO(carga_bytes), sheet_name=nh, header=None)
                header_row_idx = 0
                for idx, row in dfc_raw.iterrows():
                    row_str = " ".join([str(val).upper() for val in row.values if pd.notna(val)])
                    if "LEGAJO" in row_str or "OBRA" in row_str or "APELLIDO" in row_str:
                        header_row_idx = idx
                        break

                dfc = pd.read_excel(io.BytesIO(carga_bytes), sheet_name=nh, skiprows=header_row_idx)
                legajo_col, nombre_col, obra_col = None, None, None
                for col in dfc.columns:
                    col_str = str(col).strip().upper()
                    if "LEGAJO" in col_str: legajo_col = col
                    elif "APELLIDO" in col_str or "NOMBRE" in col_str: nombre_col = col
                    elif "OBRA" in col_str: obra_col = col

                if legajo_col is None and len(dfc.columns) > 2: legajo_col = dfc.columns[2]
                if obra_col is None: obra_col = dfc.columns[12] if len(dfc.columns) > 12 else dfc.columns[-1]

                dff = pd.read_excel(fh_fichada, sheet_name="Hoja2")

                keywords_novedades = {
                    "COMPENSATORIO": "Compensatorio", "COMPENSATORIA": "Compensatorio",
                    "ENFERMO": "Enfermo / Lic. Médica", "ENFERMEDAD": "Enfermo / Lic. Médica",
                    "ENFEMRO": "Enfermo / Lic. Médica", "ENFEMEDAD": "Enfermo / Lic. Médica",
                    "MEDICO": "Enfermo / Lic. Médica", "MÉDICO": "Enfermo / Lic. Médica",
                    "SUSPENDIDO": "Suspendido", "SUSPENSION": "Suspendido", "SUSPENSIÓN": "Suspendido",
                    "AUSENTE": "Ausente / Justificado", "AUSENCIA": "Ausente / Justificado",
                    "VACACIONES": "Vacaciones", "VACACION": "Vacaciones",
                    "SUR": "SUR / Franco Compensatorio", "LICENCIA": "Licencia / Permiso",
                }

                all_fichadas_list, fd_by_legajo, fd_by_name = [], {}, {}

                for i in range(2, len(dff)):
                    lv, nv = dff.iloc[i, 2], dff.iloc[i, 3]
                    
                    nom_str = str(nv).strip() if pd.notna(nv) else ""
                    for kw in keywords_novedades.keys():
                        nom_str = re.sub(rf"\s*-\s*{kw}.*", "", nom_str, flags=re.IGNORECASE).strip()

                    col_m_val = str(dff.iloc[i, 12]).upper() if len(dff.columns) > 12 and pd.notna(dff.iloc[i, 12]) else ""
                    row_text_sin_nombre = " ".join([str(val).upper() for idx_v, val in enumerate(dff.iloc[i].values) if idx_v != 3 and pd.notna(val)])
                    
                    novedad_detectada = None
                    for kw, label in keywords_novedades.items():
                        if kw in col_m_val or kw in row_text_sin_nombre:
                            novedad_detectada = label
                            break

                    clean_lv = limpiar_legajo(lv)
                    val_nombre_norm = normalizar_texto(nom_str)

                    info_dict = {
                        "legajo": clean_lv, "nombre": nom_str, "nombre_norm": val_nombre_norm,
                        "estado": novedad_detectada if novedad_detectada else "Presente", "es_novedad": novedad_detectada is not None,
                    }

                    if val_nombre_norm or clean_lv is not None: all_fichadas_list.append(info_dict)
                    if clean_lv is not None: fd_by_legajo[clean_lv] = info_dict
                    if val_nombre_norm: fd_by_name[val_nombre_norm] = info_dict

                def buscar_fichada(val_legajo, val_nombre):
                    clean_l = limpiar_legajo(val_legajo)
                    if clean_l is not None and clean_l in fd_by_legajo: return fd_by_legajo[clean_l], clean_l
                    val_nombre_norm = normalizar_texto(val_nombre)
                    if val_nombre_norm:
                        if val_nombre_norm in fd_by_name: return (fd_by_name[val_nombre_norm], fd_by_name[val_nombre_norm]["legajo"])
                        tokens_c = set(val_nombre_norm.split())
                        if len(tokens_c) >= 2:
                            for n_key, item in fd_by_name.items():
                                if len(tokens_c.intersection(set(n_key.split()))) >= 2:
                                    return item, item["legajo"]
                    return None, clean_l

                fichadas_emparejadas, cargados_map = set(), {}
                for _, row_c in dfc.iterrows():
                    l_val = row_c[legajo_col] if legajo_col in dfc.columns else None
                    n_val = row_c[nombre_col] if nombre_col and nombre_col in dfc.columns else None
                    info_match, clean_l = buscar_fichada(l_val, n_val)
                    o_val = str(row_c[obra_col]).strip() if obra_col and pd.notna(row_c[obra_col]) else "Sin Obra"

                    if clean_l is not None: cargados_map[clean_l] = o_val
                    if info_match:
                        if info_match["legajo"] is not None:
                            cargados_map[info_match["legajo"]] = o_val
                            fichadas_emparejadas.add(info_match["legajo"])
                        if info_match["nombre_norm"]: cargados_map[info_match["nombre_norm"]] = o_val

                no_cargados, excluidos_list, novedades_list = [], [], []
                fichadas_procesadas = set()

                for info_f in all_fichadas_list:
                    legajo_f, nombre_norm_f = info_f["legajo"], info_f["nombre_norm"]
                    f_id = legajo_f if legajo_f is not None else nombre_norm_f

                    if f_id in fichadas_procesadas or not f_id: continue
                    fichadas_procesadas.add(f_id)

                    nombre_fichada, estado_fichada, es_nov = info_f["nombre"], info_f["estado"], info_f["es_novedad"]
                    es_cargado = ((legajo_f is not None and legajo_f in cargados_map) or (legajo_f is not None and legajo_f in fichadas_emparejadas) or (nombre_norm_f and nombre_norm_f in cargados_map))

                    if legajo_f is not None and legajo_f in EXCLUIDOS_SECTOR:
                        excluidos_list.append({"LEGAJO": legajo_f, "Nombre y Apellido": (nombre_fichada or EXCLUIDOS_SECTOR[legajo_f][0]), "Sector": EXCLUIDOS_SECTOR[legajo_f][1], "Estado Fichada": estado_fichada, "Observación": "Sector no facturable"})
                    elif not es_cargado:
                        if es_nov: novedades_list.append({"LEGAJO": legajo_f if legajo_f else "SIN LEGAJO", "Nombre y Apellido": nombre_fichada, "Novedad / Estado": estado_fichada, "Observación": "Excluido por Novedad"})
                        else: no_cargados.append({"LEGAJO": legajo_f if legajo_f else "SIN LEGAJO", "Nombre y Apellido (Fichada)": nombre_fichada, "OBRA (Columna M)": "NO CARGADO EN PLANILLA", "Estado": "Fichó en planta pero NO fue cargado"})
                    else:
                        if es_nov: novedades_list.append({"LEGAJO": legajo_f if legajo_f else "SIN LEGAJO", "Nombre y Apellido": nombre_fichada, "Novedad / Estado": estado_fichada, "Observación": "Cargado registrando novedad"})

                df_no_cargados, df_excluidos, df_novedades = pd.DataFrame(no_cargados), pd.DataFrame(excluidos_list), pd.DataFrame(novedades_list)

                v1, v2, v3, v4, v5 = st.tabs(["📊 Vista General", "⚠️ Faltan Cargar (Reales)", "🚫 Sectores No Facturables", "📋 Novedades / Licencias", "🟡 Ausentes / No Ficharon"])

                with v2:
                    st.subheader("Personal que Fichó en Planta pero NO fue Cargado")
                    if not df_no_cargados.empty:
                        st.warning(f"⚠️ Se detectaron **{len(df_no_cargados)} personas** que ingresaron a trabajar y **NO figuran cargadas**.")
                        df_edit = df_no_cargados.copy()
                        df_edit["Asignar Novedad"] = "Seleccionar..."
                        edited_df = st.data_editor(df_edit, column_config={"Asignar Novedad": st.column_config.SelectboxColumn("📝 Asignar Novedad Manual", width="medium", options=["Seleccionar...", "COMPENSATORIO", "ENFERMO", "VACACIONES", "AUSENTE", "SUSPENDIDO", "SUR", "LICENCIA"])}, disabled=["LEGAJO", "Nombre y Apellido (Fichada)", "OBRA (Columna M)", "Estado"], use_container_width=True, key="editor_novedades_faltantes")

                        if st.button("☁️ Guardar Novedades Seleccionadas en Drive", type="primary"):
                            cambios = edited_df[edited_df["Asignar Novedad"].notna() & (edited_df["Asignar Novedad"] != "Seleccionar...") & (edited_df["Asignar Novedad"].astype(str).str.strip() != "")]
                            if not cambios.empty:
                                try:
                                    with st.spinner("Actualizando Excel en Google Drive..."):
                                        fh_fichada.seek(0)
                                        wb_fichada = load_workbook(fh_fichada)
                                        ws_fichada = wb_fichada["Hoja2"]

                                        for index, row_edit in cambios.iterrows():
                                            leg_buscar = limpiar_legajo(row_edit["LEGAJO"])
                                            nom_buscar = str(row_edit["Nombre y Apellido (Fichada)"]).strip()
                                            nov_asignar = str(row_edit["Asignar Novedad"]).strip()

                                            for r_idx in range(1, ws_fichada.max_row + 1):
                                                cell_val = ws_fichada.cell(row=r_idx, column=3).value
                                                cell_name = str(ws_fichada.cell(row=r_idx, column=4).value or "").strip()
                                                cell_leg = limpiar_legajo(cell_val)
                                                match_found = (cell_leg == leg_buscar) if (cell_leg is not None and leg_buscar is not None) else (nom_buscar and nom_buscar.lower() in cell_name.lower())

                                                if match_found:
                                                    ws_fichada.cell(row=r_idx, column=13).value = nov_asignar
                                                    break

                                        out_bytes = io.BytesIO()
                                        wb_fichada.save(out_bytes)
                                        out_bytes.seek(0)

                                        media_update = MediaIoBaseUpload(out_bytes, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', resumable=True)
                                        drive_service.files().update(fileId=af_selected_id, media_body=media_update).execute()

                                    st.success("¡Archivo actualizado en Google Drive con éxito!")
                                    time.sleep(2) 
                                    if hasattr(st, "rerun"): st.rerun()
                                    else: st.experimental_rerun()
                                except Exception as e:
                                    st.error(f"Error al escribir en Drive: {e}")
                            else: st.info("Elegí alguna novedad del menú antes de Guardar.")
                    else: st.success("🎉 ¡Todo el personal activo facturable está cargado!")

                with v1:
                    st.subheader("Control General de Planilla de Carga vs Fichadas")
                    def color_rows(row):
                        try:
                            l_val = row[legajo_col] if legajo_col in dfc.columns else None
                            n_val = row[nombre_col] if nombre_col and nombre_col in dfc.columns else None
                            info_match, clean_l = buscar_fichada(l_val, n_val)
                            if (clean_l is not None and clean_l in EXCLUIDOS_SECTOR) or (info_match and info_match["legajo"] in EXCLUIDOS_SECTOR): return ["background-color: #E2EFDA"] * len(row)
                            if info_match: return ["background-color: #D9E1F2"] * len(row) if info_match["es_novedad"] else ["background-color: #C6EFCE"] * len(row)
                            return ["background-color: #FFEB9C"] * len(row)
                        except Exception: return ["background-color: #FFC7CE"] * len(row)
                    st.dataframe(dfc.style.apply(color_rows, axis=1), use_container_width=True)

                with v3:
                    st.subheader("Personal de Sectores No Facturables")
                    st.dataframe(df_excluidos, use_container_width=True) if not df_excluidos.empty else st.write("Sin registros.")
                with v4:
                    st.subheader("Novedades y Licencias")
                    st.dataframe(df_novedades, use_container_width=True) if not df_novedades.empty else st.write("Sin registros.")
                with v5:
                    st.subheader("Personal Cargado que NO Registró Fichada")
                    ausentes_list = [row_c for _, row_c in dfc.iterrows() if not buscar_fichada(row_c.get(legajo_col), row_c.get(nombre_col))[0]]
                    df_ausentes = pd.DataFrame(ausentes_list) if ausentes_list else pd.DataFrame()
                    if not df_ausentes.empty:
                        st.info(f"**{len(df_ausentes)} registros** cargados sin fichada:")
                        st.dataframe(df_ausentes, use_container_width=True)
                    else: st.success("🎉 Todo el personal cargado registró su fichada.")

                wb = Workbook()
                ws1 = wb.active
                ws1.title = f"Control {nh}"
                gf, yf, rf, bf, exf = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFEB9C"), PatternFill("solid", fgColor="FFC7CE"), PatternFill("solid", fgColor="D9E1F2"), PatternFill("solid", fgColor="E2EFDA")

                for r_idx, row in enumerate(dataframe_to_rows(dfc, index=False, header=True), 1):
                    ws1.append(row)
                    if r_idx == 1: continue
                    try:
                        info_match, clean_l = buscar_fichada(row[list(dfc.columns).index(legajo_col)] if legajo_col in dfc.columns else None, row[list(dfc.columns).index(nombre_col)] if nombre_col and nombre_col in dfc.columns else None)
                        fc = exf if (clean_l is not None and clean_l in EXCLUIDOS_SECTOR) or (info_match and info_match["legajo"] in EXCLUIDOS_SECTOR) else bf if info_match and info_match["es_novedad"] else gf if info_match else yf
                        for c in range(1, len(row) + 1): ws1.cell(row=r_idx, column=c).fill = fc
                    except Exception:
                        for c in range(1, len(row) + 1): ws1.cell(row=r_idx, column=c).fill = rf

                if not df_no_cargados.empty:
                    ws2 = wb.create_sheet("Fichados Sin Carga")
                    for r_idx, row in enumerate(dataframe_to_rows(df_no_cargados, index=False, header=True), 1):
                        ws2.append(row)
                        if r_idx > 1:
                            for c in range(1, len(row) + 1): ws2.cell(row=r_idx, column=c).fill = PatternFill("solid", fgColor="FCE4D6")

                if not df_excluidos.empty:
                    ws3 = wb.create_sheet("Sectores No Facturables")
                    for r_idx, row in enumerate(dataframe_to_rows(df_excluidos, index=False, header=True), 1):
                        ws3.append(row)
                        if r_idx > 1:
                            for c in range(1, len(row) + 1): ws3.cell(row=r_idx, column=c).fill = exf

                out = io.BytesIO()
                wb.save(out)
                out.seek(0)
                st.divider()
                st.success("¡Cruce completado con éxito!")
                st.download_button("📥 Descargar Excel Formateado", data=out, file_name=f"Control_{nh}.xlsx")

            except Exception as e: st.error(f"Error al procesar: {e}")
    else: st.info("Seleccione una fecha de Google Sheets y un archivo de fichada en Drive para cruzar.")

# --- PESTAÑA: 📋 CIERRE DE PERÍODO (GESTIÓN DE OTs) ---
with t_cierre:
    st.header("📋 Notificación y Cierre de Período (Gestión de OTs)")
    if carga_bytes:
        xl_carga_cierre = pd.ExcelFile(io.BytesIO(carga_bytes))
        dias_disponibles_map = {int(m.group(1)): s for s in xl_carga_cierre.sheet_names if (m := re.search(r"(\d{1,2})[-/\.](\d{1,2})", s))}
        dias_ordenados = sorted(dias_disponibles_map.keys())

        if dias_ordenados:
            col_f1, col_f2, col_f3 = st.columns([1, 1, 1.2])
            with col_f1: 
                desde_dia_sel = st.selectbox("📅 Fecha Desde:", options=dias_ordenados, format_func=lambda d: f"Día {d:02d} ({dias_disponibles_map[d]})", index=0)
            with col_f2:
                dias_hasta_opciones = [d for d in dias_ordenados if d >= desde_dia_sel]
                hasta_dia_sel = st.selectbox("📅 Fecha Hasta:", options=dias_hasta_opciones, format_func=lambda d: f"Día {d:02d} ({dias_disponibles_map[d]})", index=len(dias_hasta_opciones)-1 if dias_hasta_opciones else 0)

            with col_f3:
                archivos_notif = obtener_archivos_notif_drive()
                mapa_notif_files = {f['name']: f for f in archivos_notif}
                opciones_notif = ["-- Sin planilla de notificaciones --"] + list(mapa_notif_files.keys())
                notif_sel_name = st.selectbox("📑 Planilla Notificaciones (Drive):", options=opciones_notif)
                
                if notif_sel_name != "-- Sin planilla de notificaciones --":
                    selected_file_info = mapa_notif_files.get(notif_sel_name, {})
                    notif_file_id = selected_file_info.get('id')
                    notif_mime_type = selected_file_info.get('mimeType')
                else:
                    notif_file_id = None
                    notif_mime_type = None

            # BOTÓN DE EJECUCIÓN DEL CRUCE
            st.write("")
            col_btn_ej, _ = st.columns([1, 3])
            with col_btn_ej:
                btn_ejecutar_cierre = st.button("🚀 Ejecutar Cruce de OTs", type="primary")

            if btn_ejecutar_cierre:
                st.session_state["cruce_cierre_ejecutado"] = True

            if st.session_state["cruce_cierre_ejecutado"]:
                with st.spinner("Procesando OTs y cruzando notificaciones desde Drive..."):
                    mapa_notif_drive, diag_info = obtener_notificaciones_drive(notif_file_id, notif_mime_type) if notif_file_id else ({}, {})

                    if mapa_notif_drive:
                        st.info(f"💡 Se detectaron **{len(mapa_notif_drive)} OTs válidas** en Drive (Columna OT: `{diag_info.get('ot_col')}`, Solicitud: `{diag_info.get('sol_col')}`, Fecha Col. D: `{diag_info.get('fecha_col')}`).")
                    elif notif_file_id:
                        st.warning("⚠️ No se pudieron asociar OTs desde la planilla seleccionada. Verificá el formato del archivo.")

                    ots_dict = {}
                    for d in dias_ordenados:
                        if d < desde_dia_sel or d > hasta_dia_sel: continue
                        s_name = dias_disponibles_map[d]
                        dfc_raw = pd.read_excel(io.BytesIO(carga_bytes), sheet_name=s_name, header=None)
                        header_row_idx = next((idx for idx, row in dfc_raw.iterrows() if any(k in " ".join([str(v).upper() for v in row.values if pd.notna(v)]) for k in ["LEGAJO", "OBRA", "APELLIDO", "OT"])), 0)
                        dfc_ot = pd.read_excel(io.BytesIO(carga_bytes), sheet_name=s_name, skiprows=header_row_idx)

                        ot_col = next((c for c in dfc_ot.columns if "OT" in str(c).upper()), None)
                        sol_col = next((c for c in dfc_ot.columns if "SOLICITUD" in str(c).upper()), None)
                        fecha_col = next((c for c in dfc_ot.columns if "FECHA" in str(c).upper()), None)
                        notif_col = next((c for c in dfc_ot.columns if "NOTIFICA" in str(c).upper()), None)
                        sector_col = next((c for c in dfc_ot.columns if any(k in str(c).upper() for k in ["SECTOR", "OBRA", "LUGAR", "AREA", "ÁREA"])), None)
                        resp_col = next((c for c in dfc_ot.columns if "RESPONSABLE" in str(c).upper()), None)

                        if ot_col:
                            for _, row_r in dfc_ot.iterrows():
                                ot_str = normalizar_ot(row_r[ot_col])
                                if not ot_str: continue

                                sol_val = str(row_r[sol_col]).strip() if sol_col and pd.notna(row_r[sol_col]) and str(row_r[sol_col]).lower() != "nan" else ""
                                
                                if fecha_col and pd.notna(row_r[fecha_col]) and str(row_r[fecha_col]).lower() != "nan":
                                    fecha_val = str(row_r[fecha_col]).strip()
                                else:
                                    fecha_val = s_name

                                notif_val = str(row_r[notif_col]).strip() if notif_col and pd.notna(row_r[notif_col]) and str(row_r[notif_col]).lower() != "nan" else ""
                                est_val = ""

                                # Cruce automático con la planilla de Drive
                                if ot_str in mapa_notif_drive:
                                    data_d = mapa_notif_drive[ot_str]
                                    if not sol_val and data_d.get("solicitud"):
                                        sol_val = data_d["solicitud"]
                                    if not notif_val and data_d.get("notificacion"):
                                        notif_val = data_d["notificacion"]
                                    if data_d.get("estado"):
                                        est_val = data_d["estado"]

                                sector_val = str(row_r[sector_col]).strip() if sector_col and pd.notna(row_r[sector_col]) and str(row_r[sector_col]).lower() != "nan" else ""
                                resp_val = str(row_r[resp_col]).strip() if resp_col and pd.notna(row_r[resp_col]) and str(row_r[resp_col]).lower() != "nan" else ""

                                if ot_str not in ots_dict:
                                    ots_dict[ot_str] = {
                                        "OT Programada": ot_str, 
                                        "N° Solicitud": sol_val, 
                                        "Estado Solicitud": est_val,
                                        "Fecha": fecha_val,
                                        "Fecha notificación final": notif_val,
                                        "Sector / Obra": sector_val,
                                        "Responsable del Trabajo": resp_val, 
                                        "Días Presente": [f"Día {d:02d}"], 
                                        "Apariciones": 1
                                    }
                                else:
                                    ots_dict[ot_str]["Apariciones"] += 1
                                    if f"Día {d:02d}" not in ots_dict[ot_str]["Días Presente"]: ots_dict[ot_str]["Días Presente"].append(f"Día {d:02d}")
                                    if not ots_dict[ot_str]["N° Solicitud"] and sol_val: ots_dict[ot_str]["N° Solicitud"] = sol_val
                                    if not ots_dict[ot_str]["Estado Solicitud"] and est_val: ots_dict[ot_str]["Estado Solicitud"] = est_val
                                    if not ots_dict[ot_str]["Fecha"] and fecha_val: ots_dict[ot_str]["Fecha"] = fecha_val
                                    if not ots_dict[ot_str]["Fecha notificación final"] and notif_val: ots_dict[ot_str]["Fecha notificación final"] = notif_val
                                    if not ots_dict[ot_str]["Sector / Obra"] and sector_val: ots_dict[ot_str]["Sector / Obra"] = sector_val
                                    if not ots_dict[ot_str]["Responsable del Trabajo"] and resp_val: ots_dict[ot_str]["Responsable del Trabajo"] = resp_val

                    df_ots_unicas = pd.DataFrame(list(ots_dict.values()))

                if not df_ots_unicas.empty:
                    df_ots_unicas["Días Presente"] = df_ots_unicas["Días Presente"].apply(lambda x: ", ".join(x))
                    
                    # Cálculo dinámico de métricas / contadores
                    tot_ots = len(df_ots_unicas)
                    con_solic = len(df_ots_unicas[df_ots_unicas["N° Solicitud"].astype(str).str.strip() != ""])
                    
                    es_aprobada = df_ots_unicas.apply(
                        lambda r: bool(
                            re.search(r'APROB|LIBER|NOTIF|FINAL', str(r["Estado Solicitud"]).upper()) or 
                            (pd.notna(r["Fecha notificación final"]) and str(r["Fecha notificación final"]).strip() != "")
                        ), axis=1
                    )
                    
                    ots_aprobadas = int(es_aprobada.sum())
                    ots_pendientes = tot_ots - ots_aprobadas

                    k_c1, k_c2, k_c3, k_c4 = st.columns(4)
                    k_c1.metric("📌 Total OTs Únicas", tot_ots)
                    k_c2.metric("📋 OTs con Solicitud", con_solic)
                    k_c3.metric("🟢 OTs Aprobadas / Notificadas", ots_aprobadas)
                    k_c4.metric("⏳ OTs Pendientes", ots_pendientes)
                    
                    st.divider()
                    st.write("📝 **Completá o modificá los datos directamente en la tabla:**")
                    
                    edited_ots_df = st.data_editor(
                        df_ots_unicas, 
                        column_config={
                            "OT Programada": st.column_config.TextColumn("OT Programada", disabled=True), 
                            "N° Solicitud": st.column_config.TextColumn("N° Solicitud"),
                            "Estado Solicitud": st.column_config.TextColumn("Estado Solicitud"),
                            "Fecha": st.column_config.TextColumn("Fecha"),
                            "Fecha notificación final": st.column_config.TextColumn("Fecha notificación final"),
                            "Sector / Obra": st.column_config.TextColumn("Sector / Obra"),
                            "Responsable del Trabajo": st.column_config.TextColumn("Responsable del Trabajo"), 
                            "Días Presente": st.column_config.TextColumn("Días Presente", disabled=True), 
                            "Apariciones": st.column_config.NumberColumn("Apariciones", disabled=True)
                        }, 
                        use_container_width=True, 
                        key="editor_ots_cierre"
                    )

                    if st.button("💾 Guardar y Sincronizar Google Sheets", type="primary"):
                        try:
                            cambios_ots = {
                                str(r["OT Programada"]).strip(): {
                                    "Solicitud": str(r["N° Solicitud"]).strip() if pd.notna(r["N° Solicitud"]) and str(r["N° Solicitud"]).lower() != "nan" else "", 
                                    "Estado": str(r["Estado Solicitud"]).strip() if pd.notna(r["Estado Solicitud"]) and str(r["Estado Solicitud"]).lower() != "nan" else "",
                                    "Fecha": str(r["Fecha"]).strip() if pd.notna(r["Fecha"]) and str(r["Fecha"]).lower() != "nan" else "",
                                    "FechaNotificacion": str(r["Fecha notificación final"]).strip() if pd.notna(r["Fecha notificación final"]) and str(r["Fecha notificación final"]).lower() != "nan" else "",
                                    "Sector": str(r["Sector / Obra"]).strip() if pd.notna(r["Sector / Obra"]) and str(r["Sector / Obra"]).lower() != "nan" else "",
                                    "Responsable": str(r["Responsable del Trabajo"]).strip() if pd.notna(r["Responsable del Trabajo"]) and str(r["Responsable del Trabajo"]).lower() != "nan" else ""
                                } for _, r in edited_ots_df.iterrows()
                            }
                            with st.spinner("Sincronizando con Google Sheets..."):
                                res = requests.post(WEBAPP_URL, json={"desdeDia": desde_dia_sel, "hastaDia": hasta_dia_sel, "cambios": cambios_ots})
                            if res.status_code == 200:
                                st.success("🟢 ¡Sincronizado con éxito!")
                            else: st.error(f"Error {res.status_code}")
                        except Exception as e: st.error(f"Error al enviar datos: {e}")
                else: st.warning("No se encontraron OTs cargadas.")
        else: st.info("No se detectaron días cargados en la planilla.")
    else: st.info("Conectando con Google Sheets...")

with t_quincena:
    st.header("⏱️ Recuento y Validación Quincenal de Horas")
    if carga_bytes:
        col_q1, col_q2 = st.columns([2, 1])
        with col_q1: quincena_sel = st.radio("Seleccionar Período:", options=["1ª Quincena (Días 1 al 15)", "2ª Quincena (Días 16 al 31)", "Todo el Mes"], horizontal=True)
        with col_q2: btn_calc_q = st.button("📊 Calcular Resumen Quincenal", type="primary")

        if btn_calc_q or st.session_state.get("ver_quincena", False):
            st.session_state["ver_quincena"] = True
            try:
                xl_carga = pd.ExcelFile(io.BytesIO(carga_bytes))
                registros_quincena = []

                for sheet in xl_carga.sheet_names:
                    match = re.search(r"(\d{1,2})[-/\.](\d{1,2})", sheet)
                    if not match: continue
                    dia, mes = int(match.group(1)), int(match.group(2))
                    if quincena_sel == "1ª Quincena (Días 1 al 15)" and dia > 15: continue
                    elif quincena_sel == "2ª Quincena (Días 16 al 31)" and dia <= 15: continue

                    dfc_raw = pd.read_excel(io.BytesIO(carga_bytes), sheet_name=sheet, header=None)
                    header_row_idx = next((idx for idx, row in dfc_raw.iterrows() if any(k in " ".join([str(v).upper() for v in row.values if pd.notna(v)]) for k in ["LEGAJO", "OBRA", "APELLIDO"])), 0)
                    dfc_q = pd.read_excel(io.BytesIO(carga_bytes), sheet_name=sheet, skiprows=header_row_idx)

                    legajo_col = next((c for c in dfc_q.columns if "LEGAJO" in str(c).upper()), None)
                    nombre_col = next((c for c in dfc_q.columns if any(k in str(c).upper() for k in ["APELLIDO", "NOMBRE", "PERSONAL"])), None)
                    trabajo_col = next((c for c in dfc_q.columns if "TRABAJO" in str(c).upper()), None)
                    tipo_col = next((c for c in dfc_q.columns if any(k in str(c).upper() for k in ["TIPO DE HORA", "TIPO"])), None)

                    if legajo_col and trabajo_col and tipo_col:
                        for _, row in dfc_q.iterrows():
                            leg_clean = limpiar_legajo(row[legajo_col])
                            nom = str(row[nombre_col]).strip() if nombre_col and pd.notna(row[nombre_col]) else (f"Legajo {leg_clean}" if leg_clean else "Sin Nombre")
                            hs = extraer_horas(row[trabajo_col])
                            tipo_raw = str(row[tipo_col]).strip().upper() if pd.notna(row[tipo_col]) else "HS. NORMALES"
                            tipo_cat = "Hs. 50%" if "0.5" in tipo_raw or "50" in tipo_raw else "Hs. 100%" if "1" in tipo_raw or "100" in tipo_raw else "Hs. Normales"

                            if leg_clean is not None or nom != "Sin Nombre":
                                registros_quincena.append({"Legajo": leg_clean if leg_clean else 0, "Nombre": nom, "Horas": hs, "Tipo": tipo_cat})

                df_q_all = pd.DataFrame(registros_quincena)
                if not df_q_all.empty:
                    piv_q = df_q_all.pivot_table(index=["Legajo", "Nombre"], columns="Tipo", values="Horas", aggfunc="sum", fill_value=0.0).reset_index()
                    for col_h in ["Hs. Normales", "Hs. 50%", "Hs. 100%"]:
                        if col_h not in piv_q.columns: piv_q[col_h] = 0.0
                    piv_q["Total Hs. Cargadas"] = piv_q["Hs. Normales"] + piv_q["Hs. 50%"] + piv_q["Hs. 100%"]

                    fichadas_horas_dict = {}
                    with st.spinner("Analizando fichadas en Google Drive..."):
                        archivos_drive = obtener_archivos_drive()
                        mapa_archivos_q = {f['name']: f['id'] for f in archivos_drive if f['name'].endswith(('.xlsx', '.xltx'))}
                        
                        for f_loc, f_id in mapa_archivos_q.items():
                            match_f = re.search(r"Fichada_(\d{1,2})-(\d{1,2})", f_loc)
                            if match_f:
                                d_f = int(match_f.group(1))
                                if quincena_sel == "1ª Quincena (Días 1 al 15)" and d_f > 15: continue
                                elif quincena_sel == "2ª Quincena (Días 16 al 31)" and d_f <= 15: continue
                                try:
                                    fh_loc = descargar_archivo_drive(f_id)
                                    dff_loc = pd.read_excel(fh_loc, sheet_name="Hoja2")
                                    for idx_fl in range(2, len(dff_loc)):
                                        l_f_clean = limpiar_legajo(dff_loc.iloc[idx_fl, 2])
                                        if l_f_clean is not None:
                                            fichadas_horas_dict[l_f_clean] = fichadas_horas_dict.get(l_f_clean, 0.0) + calcular_horas_fichada(dff_loc.iloc[idx_fl, 4], dff_loc.iloc[idx_fl, 5])
                                except Exception: pass

                    df_q_final = pd.merge(piv_q, pd.DataFrame(list(fichadas_horas_dict.items()), columns=["Legajo", "Total Hs. Fichadas (Reloj)"]), on="Legajo", how="left").fillna(0.0)
                    df_q_final["Diferencia (Cargado vs Fichado)"] = df_q_final["Total Hs. Cargadas"] - df_q_final["Total Hs. Fichadas (Reloj)"]

                    kpi1, kpi2, kpi3, kpi4 = st.columns(4)
                    kpi1.metric("📊 Total Hs. Normales", f"{df_q_final['Hs. Normales'].sum():.1f} hs")
                    kpi2.metric("⚡ Total Hs. 50%", f"{df_q_final['Hs. 50%'].sum():.1f} hs")
                    kpi3.metric("🔥 Total Hs. 100%", f"{df_q_final['Hs. 100%'].sum():.1f} hs")
                    kpi4.metric("⏱️ Total Horas Cargadas", f"{df_q_final['Total Hs. Cargadas'].sum():.1f} hs")
                    
                    st.divider()
                    st.dataframe(df_q_final[["Legajo", "Nombre", "Hs. Normales", "Hs. 50%", "Hs. 100%", "Total Hs. Cargadas", "Total Hs. Fichadas (Reloj)", "Diferencia (Cargado vs Fichado)"]].style.applymap(lambda val: "background-color: #FFEB9C; color: #856404;" if val > 0 else "background-color: #F8D7DA; color: #721C24;" if val < 0 else "background-color: #C6EFCE; color: #155724;", subset=["Diferencia (Cargado vs Fichado)"]), use_container_width=True)

                    out_q = io.BytesIO()
                    with pd.ExcelWriter(out_q, engine="openpyxl") as writer: df_q_final.to_excel(writer, index=False, sheet_name="Resumen Quincenal")
                    out_q.seek(0)
                    st.download_button("📥 Descargar Resumen Quincenal (Excel)", data=out_q, file_name=f"Resumen_Horas_{quincena_sel.replace(' ', '_')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                else: st.warning("No se encontraron registros de horas cargadas.")
            except Exception as e_q: st.error(f"Error al calcular: {e_q}")
    else: st.info("Conectando con Google Sheets...")
